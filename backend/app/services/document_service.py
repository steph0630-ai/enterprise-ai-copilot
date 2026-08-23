"""文档服务：处理文件上传的业务逻辑（保存文件 + 写数据库 + 多格式解析 + 切分 + 异步入库）

Day 18 大改造：从"上传时同步入库"变成"上传秒回，后台线程慢慢入库"。
- 大文件（几十 MB 到 200MB）不再把上传请求挂死：请求只负责收文件、建记录，
  真正的解析/向量化挪到 ingest_document_job（FastAPI BackgroundTasks 后台跑）。
- 保存改成流式写盘：边写边数大小，超限即中止，不整文件读进内存。

Day 20 多格式：解析从"只认 PDF"升级成"按扩展名分发的解析层"——
parse_document 把 .pdf/.docx/.txt/.md 分发给对应解析器，全部返回 list[str]，
对下游 split_text 完全透明。加格式 = 写一个解析函数 + 分发表加一行。
顺手修：路径穿越（sanitize_filename）、空文档显式失败（不再静默 processed+0 chunk）。
"""

import io
import logging
import os
import re
import threading
from collections import Counter
from pathlib import Path

from fastapi import HTTPException

logger = logging.getLogger(__name__)  # Day 23：丢图可感知（日志在 main.py 配了 basicConfig）
from langchain_text_splitters import RecursiveCharacterTextSplitter
import pdfplumber
# Day 20：python-docx 解析 .docx。
# 关键坑：docx.Document 是【工厂函数】不是类——打开文件用它（返回 DocumentPart.document 实例），
# 但 isinstance 判断必须用真正的类 docx.document.Document，两者是同一个名字、不同类型。
# （Document 名字已被下面的业务模型占用，所以真正的类起别名 DocxDocumentType。）
import docx
from docx.document import Document as DocxDocumentType
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from sqlalchemy.orm import Session

from app.ai.embedding import embedding_service
from app.ai.vision_service import vision_service
from app.core.config import settings
from app.database.session import SessionLocal
from app.models.document import Document
from app.models.document_chunk import DocumentChunk
from app.vectorstore.store import vector_store

# 上传文件的存放目录（相对 backend 根目录）
# 用 Path 对象，自动处理 Windows/Linux 的路径分隔符差异
STORAGE_DIR = Path(__file__).resolve().parent.parent / "storage" / "files"

# 并发保险（Day 18）：向量库 + chunk 表是本地快写，用锁串起来，
# 防止两个上传同时入库时 Chroma/sqlite 撞车。
# 注意：embedding 的 API 调用（慢、网络）放在锁外面，并发不排队。
_ingest_lock = threading.Lock()


def sanitize_filename(filename: str | None) -> str:
    """剥掉路径成分，防路径穿越：'..\\..\\x.pdf' / '../../x.pdf' -> 'x.pdf'

    必须先把反斜杠统一成斜杠再交给 Path：Linux 容器里 Path 不把 '\\' 当分隔符，
    直接 Path('..\\..\\x.pdf').name 会原样返回整串，穿越就堵不住了。
    Path(...).name 只留最后一段，'../' 全部失效。
    """
    if not filename:
        return ""
    return Path(filename.replace("\\", "/")).name


def get_extension(filename: str | None) -> str:
    """返回规范化后的小写扩展名（含点）；空名/无扩展名返回 ''

    例：'A.PDF' -> '.pdf'，'report.txt' -> '.txt'，'README' -> ''
    """
    if not filename:
        return ""
    return Path(filename.strip().replace("\\", "/")).suffix.lower()


def save_uploaded_file(src, filename: str, max_size: int | None = None) -> str:
    """把上传的文件流写到磁盘，返回文件路径（相对 backend 根目录）

    Day 18 改成流式写入：不把整文件读进内存，1MB 一块边写边数字节，
    超上限就删掉半截文件并报 413——企业大文件不能靠"整读进内存"硬撑。

    参数：
        src: 文件对象（Starlette 的 UploadFile.file，读取前已 seek 回开头）
        filename: 原始文件名，如 "合同.pdf"
        max_size: 允许的最大字节数，默认取 settings.MAX_DOC_SIZE（200MB）
    """
    if max_size is None:
        max_size = settings.MAX_DOC_SIZE

    STORAGE_DIR.mkdir(parents=True, exist_ok=True)

    # Day 20 修路径穿越：原实现直接 STORAGE_DIR / filename，filename 含 '../'
    # 可把文件写到 storage 目录之外。现在剥掉路径成分，只留文件名。
    safe_name = sanitize_filename(filename)
    if not safe_name:
        raise HTTPException(status_code=400, detail="文件名无效")

    # 拼出完整路径，然后分块拷贝
    file_path = STORAGE_DIR / safe_name
    written = 0
    with open(file_path, "wb") as out:
        while True:
            chunk = src.read(1024 * 1024)  # 每次最多 1MB
            if not chunk:  # 读到头了
                break
            written += len(chunk)
            if written > max_size:  # 超限：删半截文件，别留垃圾
                out.close()
                file_path.unlink(missing_ok=True)
                raise HTTPException(
                    status_code=413,
                    detail=f"文件超过大小上限（{max_size // 1024 // 1024}MB）",
                )
            out.write(chunk)

    # 返回相对路径（相对于 backend 根目录）
    # 用 os.path.relpath 计算，保证可移植
    backend_root = Path(__file__).resolve().parent.parent.parent
    return os.path.relpath(file_path, backend_root)


def create_document_record(
    db: Session,
    filename: str,
    file_path: str,
    knowledge_base_id: int,
    status: str = "uploaded",
) -> Document:
    """往 documents 表插一条记录，返回这个 Document 对象

    Day 18：加 status 参数——上传入口走 "uploading"（还没开始入库），
    老调用不传则保持默认 "uploaded"。
    """
    doc = Document(
        filename=filename,
        file_path=file_path,
        knowledge_base_id=knowledge_base_id,
        status=status,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)  # 让 doc 拿到数据库生成的自增 id
    return doc


def parse_pdf(file_path: str) -> list[str]:
    """解析 PDF，返回每一页的文本列表

    为什么用 pdfplumber 而不是 pypdf（Day 18 踩坑）：
    某些 PDF（尤其从 WPS/在线转换器生成的中文 PDF）里，字形→文字的映射表
    ToUnicode 是坏的。pypdf 照表抽，抽出来全是乱码；眼睛看得见、Ctrl+F 搜不到。
    pdfplumber 底层是 pdfminer，会自己重建字形编码，能把这种 PDF 救回来。

    为什么表格要转 Markdown（Day 25.3）：
    年报"未抵销的递延所得税负债"这类财务表是"科目 × 多期 × 多列"，只 extract_text()
    会把列结构打平——"应纳税暂时性差异/递延所得税负债/期末/期初"这些列头和数据失去
    对应，模型把"差异"列的数字当成"负债"列，答案数字就错（真实年报测出来的坑）。
    Markdown 表格保留"列头↔数据"的对应关系，模型才读得懂列。

    参数：
        file_path: PDF 文件的路径

    返回：
        ["第一页文本", "第二页文本", ...]
    """
    with pdfplumber.open(file_path) as pdf:
        pages = [_page_to_markdown(page) for page in pdf.pages]
    # Day 25.3：剥离重复页眉/页脚 + 页码（真实年报页脚"103/155"污染 chunk）
    return _strip_headers_footers(pages)


def _page_to_markdown(page) -> str:
    """一页 → 正文 + Markdown 表格；表格区域不再从正文重复抽取。"""
    table_objects = list(page.find_tables()) if hasattr(page, "find_tables") else []
    if table_objects:
        tables = [table.extract() for table in table_objects if table.extract()]
        text = _extract_text_outside_tables(page, [table.bbox for table in table_objects])
    else:
        # 兼容测试桩和旧版 pdfplumber：没有 bbox 时保留原行为。
        text = page.extract_text() or ""
        tables = [t for t in (page.extract_tables() or []) if t]
    if not tables:
        return text
    md = "\n\n".join(
        "[表格]\n" + _table_to_markdown(t) + "\n[表格结束]" for t in tables
    )
    return text + "\n\n" + md


def _extract_text_outside_tables(page, table_bboxes: list[tuple]) -> str:
    """提取表格外文字，避免 extract_text 与 Markdown 表格重复。"""
    if not table_bboxes or not hasattr(page, "filter"):
        return page.extract_text() or ""

    def keep(obj: dict) -> bool:
        if obj.get("object_type") != "char":
            return True
        x = (obj.get("x0", 0) + obj.get("x1", 0)) / 2
        y = (obj.get("top", 0) + obj.get("bottom", 0)) / 2
        return not any(x0 <= x <= x1 and top <= y <= bottom for x0, top, x1, bottom in table_bboxes)

    return page.filter(keep).extract_text() or ""


def _should_join_without_space(left: str, right: str) -> bool:
    """相邻两段应无缝拼接（Day 25.4）：数字或中文被 PDF 折行劈开时不能留空格。

    两种情况：
    ① 数字被折行（835,160,639.35 显示成 835,160,639.3\\n5）→ 接回一个数。
    ② 中文词被折行（发出商品显示成 发出商\\n品）→ 中文词内本没有空格，
       若拼成"发出商 品"，"发出商品"被拆成两词，和"库存商 品"几乎一样，
       模型会认错行（真实故障：问"发出商品"答成"库存商品"的数）。
    """
    # ① 数字折行：left 已有小数点、以数字/点结尾，right 是 1~2 位纯数字
    if "." in left and re.search(r"[0-9.]$", left) and re.fullmatch(r"\d{1,2}", right):
        return True
    # ② 中文折行：left 末字和 right 首字都是汉字 → 无空格连上（中文句内无词间空格）
    if re.search(r"[一-鿿]", left[-1:]) and re.match(r"[一-鿿]", right[:1]):
        return True
    return False


def _flatten_cell(c) -> str:
    """一个单元格 → 单行文本（Day 25.4：修 PDF 折行把数字/中文词劈开）

    原来直接 " ".join(str(c).split())：pdfplumber 抽出的单元格里，长金额会被 PDF
    折行成多行（835,160,639.35 → 835,160,639.3\\n5；发出商品 → 发出商\\n品），
    压平后换行变空格——数字被劈成"835,160,639.3 5"（读作缺一位），
    中文词被拆成"发出商 品"（和"库存商 品"几乎一样，模型认错行）。
    这里按行拼接，只在"数字折行 / 中文折行"处无缝接回，其余换行用单个空格。
    """
    text = "" if not c else str(c)
    lines = [ln.strip() for ln in text.split("\n")]
    parts = [ln for ln in lines if ln]
    out = ""
    for part in parts:
        if out and _should_join_without_space(out, part):
            out += part  # 数字 / 中文词被折行劈开 → 无缝接回
        else:
            out = f"{out} {part}" if out else part
    return " ".join(out.split())


def _table_to_markdown(table: list) -> str:
    """pdfplumber 抽出的表格（行×列 嵌套 list）→ Markdown 表格

    空单元格 → 空串；单元格内多行 → 空格压平（避免 \n 把 Markdown 行撑爆）。
    Day 25.3 表头增强：财务表常有两行表头（第一行"项目|期末余额||期初余额|"，
    第二行才是"可抵扣暂时性差异|递延所得税资产"），只取第一行会把真正列名
    当数据行，模型读错列。检测到第二行也是表头就合并成一行。
    """
    rows = []
    for row in table:
        cells = []
        for c in row:
            cells.append(_flatten_cell(c))
        rows.append(cells)
    if not rows:
        return ""
    # 表头行数：第二行像表头（无金额、含列名词）→ 合并两行当表头
    header_rows = 2 if len(rows) >= 2 and _looks_like_header_row(rows[1]) else 1
    header = _merge_header_rows(rows[:header_rows])
    body = rows[header_rows:]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * len(header)) + " |",
    ]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


def _looks_like_header_row(cells: list) -> bool:
    """这一行像是表头：不含金额（有逗号的数字），且含表头特征词

    财务表数据行一定带金额（1,417,003.17），表头行只有列名。
    用"无金额 + 含表头词"做保守判断，避免把数据行误当表头。
    """
    joined = " ".join(cells)
    if re.search(r"\d{1,3},\d{3}", joined):
        return False
    return any(kw in joined for kw in (
        "项目", "余额", "差异", "资产", "负债", "期末", "期初", "年份", "数量",
    ))


def _merge_header_rows(rows: list[list]) -> list[str]:
    """把多行表头逐列拼成一行：同列的"期末余额"+"递延所得税负债" → "期末余额 递延所得税负债"

    关键：先做"向右填充"。跨列合并的表头（第一行"期末余额"横跨两列）在
    pdfplumber 里只在第一列有值、第二列是空，不填充的话第二列就丢了
    "期末余额"这个前缀，只剩"递延所得税负债"，期末/期初区分就没了。
    """
    n = max(len(r) for r in rows)
    # 1. 每行向右填充（跨列合并：本列空就继承上一非空列的值）
    filled = []
    for r in rows:
        last = ""
        row = []
        for j in range(n):
            v = r[j] if j < len(r) else ""
            if v:
                last = v
            row.append(last)
        filled.append(row)
    # 2. 逐列拼接各行
    merged = []
    for j in range(n):
        parts = [filled[i][j] for i in range(len(filled))]
        merged.append(" ".join(p for p in parts if p))
    return merged


# ========== Day 25.3：页眉页脚剥离 ==========
# 真实年报的页脚"103/155"、每页重复的"XX公司 2026年半年度报告"混进 chunk，
# 污染检索结果。两种手法组合：
#   ① 跨页重复行：公司名/文档标题/机密水印每页都出现 → 重复检测（≥ 60% 页数）
#   ② 页码模式：每页变化的页码（103/155、第5页、- 5 -）→ 正则识别
_HEADER_FOOTER_FREQ = 0.6   # 出现在 ≥ 60% 页数 = 重复页眉/页脚
_MIN_STRIP_PAGES = 3        # 少于 3 页不做（单页没有"重复"意义）
# 页码模式：只认"像页码"的行，不认裸数字（表格打平出的"1417003.17"绝不能误删）
_PAGE_NUM_RE = re.compile(
    r"^\s*(?:第\s*\d+\s*页(?:\s*共\s*\d+\s*页)?"
    r"|\d{1,4}\s*/\s*\d{1,4}"
    r"|[-—–]\s*\d+\s*[-—–])\s*$"
)
# 纯数字/金额/日期样行（数字+标点）：不因"跨页重复"被剥，防误删表格数字
_NUMERIC_LINE_RE = re.compile(r"^[\d\s.,，()%¥￥$€\-—–]+$")
# 表格包裹标记（Day 25.4）：页眉页脚剥离绝不能把它当"重复行"删掉。
# 真实年报几乎全是表格，[表格] 出现在绝大多数页 → 会命中"≥60% 重复短行"。
# 一旦被删，所有表失去"整表一块"识别，split 会把表格当正文按 500 字劈开，
# 出现"库存商品、发出商品被劈进两个 chunk、模型答错行"的故障。
_TABLE_MARKERS = {"[表格]", "[表格结束]"}


def _strip_headers_footers(pages: list[str]) -> list[str]:
    """剥离每页重复的页眉/页脚行和页码，返回剥后的页列表

    为什么在纯文本层做而不是坐标层：extract_text 已经把页眉/页脚/正文混成行，
    "跨页重复 + 页码模式"两个特征足够稳，且纯函数好测。
    保守原则：宁可漏剥（如纯数字页码），不能误删正文内容。
    """
    if len(pages) < _MIN_STRIP_PAGES:
        return pages
    # 1. 统计每行出现在"几页"（每页只计一次，防一页内重复干扰）
    line_pages: Counter = Counter()
    for text in pages:
        for line in {ln.strip() for ln in text.split("\n") if ln.strip()}:
            line_pages[line] += 1
    # 2. 页眉/页脚 = 重复出现（≥ 60% 页数）的短文本行 + 页码模式行
    threshold = len(pages) * _HEADER_FOOTER_FREQ
    noise = {
        ln for ln, cnt in line_pages.items()
        if cnt >= threshold and len(ln) <= 50
        and not _NUMERIC_LINE_RE.match(ln)  # 纯数字行（表格金额等）不因重复误删
        and ln not in _TABLE_MARKERS        # 表格标记绝不算页眉页脚（见 Day 25.4）
    }
    for text in pages:
        for line in text.split("\n"):
            s = line.strip()
            if s and _PAGE_NUM_RE.match(s):
                noise.add(s)
    # 3. 逐页剔除
    return [
        "\n".join(ln for ln in text.split("\n") if ln.strip() not in noise)
        for text in pages
    ]


def parse_txt(file_path: str) -> list[str]:
    """解析 .txt / .md：整篇文本作为一"页"返回（对 split_text 透明）

    编码坑：企业 Windows 上传的 txt 常是 GBK，而 Docker 容器默认 UTF-8，
    直接按 UTF-8 读会 UnicodeDecodeError 崩掉。策略是"读字节一次、逐级试解码"：
      - utf-8-sig：等于 utf-8 + 自动剥 BOM，省一个编码位
      - gbk：救 Windows 老文件（GBK 是 GB2312 超集，中文系统通用）
      - latin-1：永不抛错（256 种字节全映射成字符），保证坏编码也能读出不崩
    只读一次磁盘（open+read），而不是每个编码各 open 一次——大文件读 2~3 遍太浪费。
    """
    with open(file_path, "rb") as f:
        raw = f.read()

    for encoding in ("utf-8-sig", "gbk", "latin-1"):
        try:
            return [raw.decode(encoding)]
        except UnicodeDecodeError:
            continue

    # 理论到不了：latin-1 永远不会抛 UnicodeDecodeError
    raise ValueError("无法解码文本文件（不支持的编码）")


def _iter_block_items(parent):
    """按文档 body 顺序产出段落和表格；parent 是文档对象或单元格

    为什么不用 doc.paragraphs / doc.tables：它们只给"顶层段落"和"顶层表格"，
    两者都丢掉了"段落在表格前还是后"的顺序。制度类文档常是"条款段落 + 中间
    嵌一张表"，按顺序抽出来接成一段，chunk 语义才连贯。
    """
    parent_elm = parent.element.body if isinstance(parent, DocxDocumentType) else parent._tc
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield DocxParagraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield DocxTable(child, parent)


def _cell_text(cell) -> str:
    """单元格文本：段落 + 嵌套表格都要（制度类文档常有嵌套表）"""
    parts = [p.text for p in cell.paragraphs if p.text.strip()]
    for t in cell.tables:
        parts.append(_table_text(t))
    return "\n".join(parts)


def _table_text(table) -> str:
    """docx 表格 → Markdown 表格（Day 25.3：和 PDF 表格统一，保留列头↔数据）

    原来用 " | " 分隔打平，多列表格（财务表）的列语义会丢——模型分不清
    "差异/负债""期末/期初"哪列对哪列，数字就读错。Markdown 带表头行 +
    分隔行，列头↔数据的对应关系保留（和 _table_to_markdown 同一套）。
    """
    rows = []
    for row in table.rows:
        rows.append([_cell_text(c).replace("\n", " ") for c in row.cells])
    return _table_to_markdown(rows)


def parse_docx(file_path: str) -> list[str]:
    """解析 .docx：正文段落 + 表格按文档顺序抽，整篇作为一"页"返回

    表头/页脚明确跳过：它们每页重复（页码/公司名/文档标题），混进 chunk 是纯噪声。
    损坏文件（或伪装成 .docx 的假文件）python-docx 会抛 BadZipFile 之类，
    包一层转成带中文说明的 ValueError，让 error_message 是"看得懂的话"。
    """
    try:
        document = docx.Document(file_path)  # 工厂函数打开文件，返回 Document 实例
    except Exception as e:
        raise ValueError(f"无法解析 Word 文档（文件可能损坏或不是真正的 .docx）：{e}")

    blocks: list[str] = []
    for block in _iter_block_items(document):
        if isinstance(block, DocxParagraph):
            text = block.text.strip()
            if text:
                blocks.append(text)
        elif isinstance(block, DocxTable):
            text = _table_text(block).strip()
            if text:
                blocks.append(text)

    return ["\n\n".join(blocks)]


# ========== 图片理解（Day 21：混图文档的"看图说话"） ==========
# 图片理解是【增强层】不是【必要层】：图 → 多模态模型 → 文字描述 → 拼回文本流，
# 下游 split_text/embedding 完全透明。任何一张图失败都静默跳过，不拖垮整个文档入库。
# 小图（logo/装饰图标）跳过——它们宽或高 < MIN_IMAGE_SIZE，是噪声不是内容。
MIN_IMAGE_SIZE = 100


def extract_pdf_images(file_path) -> list[tuple[int, bytes, str]]:
    """从 PDF 抠出值得描述的图片，返回 [(页索引, PNG 字节, mime), ...]

    为什么不用 PyMuPDF：它是 PDF 抠图的标准工具，但它是 AGPL 授权——企业项目
    不开源就有合规隐患。pdfplumber 自带渲染能力：page.images 给出每张图的 bbox，
    page.crop(bbox).to_image() 把该区域渲染成位图，够 VL 模型看，还少一个重依赖。
    """
    images: list[tuple[int, bytes, str]] = []
    with pdfplumber.open(file_path) as pdf:
        for page_idx, page in enumerate(pdf.pages):
            for img in page.images:
                # srcsize 是 (宽, 高) 像素；缺失时放行（宁可多看，别漏内容）
                srcsize = img.get("srcsize")
                w, h = srcsize or (0, 0)
                if w and h and (w < MIN_IMAGE_SIZE or h < MIN_IMAGE_SIZE):
                    continue  # 小图：logo/装饰图标，跳过
                bbox = (img["x0"], img["top"], img["x1"], img["bottom"])
                rendered = page.crop(bbox).to_image(resolution=150).original
                buf = io.BytesIO()
                rendered.save(buf, format="PNG")
                images.append((page_idx, buf.getvalue(), "image/png"))
    return images


def extract_docx_images(file_path) -> list[tuple[int, bytes, str]]:
    """从 .docx 抠出值得描述的图片，返回 [(块索引, 图片字节, mime), ...]

    python-docx 没有"图片列表"这个 API，图藏在 XML 的 a:blip 里：
      - document.element.body.iter(qn('a:blip'))：遍历所有图片引用
      - blip.get(qn('r:embed'))：拿到 relationship id（rIdN）
      - document.part.related_parts[rid]：取到 ImagePart（.blob 字节 / .content_type mime）
    part.image.px_width/px_height 是真实像素尺寸，用来过滤小图。

    parse_docx 把整篇合并成一页返回，所以这里块索引统一记 0（描述拼到那页末尾）。
    """
    images: list[tuple[int, bytes, str]] = []
    document = docx.Document(file_path)
    for blip in document.element.body.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue  # 有 blip 没 embed（外链图），跳过
        try:
            part = document.part.related_parts[rid]
        except KeyError:
            continue  # relationship 引用悬空（损坏文档），跳过不崩
        if part.image.px_width < MIN_IMAGE_SIZE or part.image.px_height < MIN_IMAGE_SIZE:
            continue  # 小图跳过
        images.append((0, part.blob, part.content_type))
    return images


def enrich_pages_with_images(filename: str, pages_text: list[str]) -> tuple[list[str], dict]:
    """图片理解增强层：文档里的图 → 多模态模型描述 → 拼回对应页文本末尾

    返回 (pages_text, stats)：stats = {"total": 尝试处理的图数, "success": 成功,
    "failed": 失败跳过}，且恒有 total == success + failed（调用方 _ingest 打总结日志）。

    快速路径（txt/md 没有内嵌图）直接原样返回，零开销。失败降级：任何一张图
    （VL 挂/超时/图片损坏）跳过，文档照常入库——图片理解是锦上添花，
    不该因为看图失败就让整个文档 failed。这和 Day 20 空文档显式失败不冲突：
    那是核心能力（没字可抽 = 文档没用）必须报，这是可降级能力，降级后仍可用。

    Day 23 教训：降级可以静默，但不能无声。之前失败直接 continue，丢图无人知晓
    （实测尚硅谷 PDF 6 张图只进库 4 条，2 张被静默丢了）。现在失败打 warning 日志
    + 计数，让"这张图没进去"有迹可循。

    数量风控：最多处理 settings.VISION_MAX_IMAGES 张，防一张 100 图的 PPT
    批量上传打爆 API 账单。
    """
    empty_stats = {"total": 0, "success": 0, "failed": 0}
    if not pages_text:
        return pages_text, empty_stats

    # 快速路径：只有 pdf/docx 有内嵌图，其余原样返回
    ext = get_extension(filename)
    if ext == ".pdf":
        images = extract_pdf_images(STORAGE_DIR / filename)
    elif ext == ".docx":
        images = extract_docx_images(STORAGE_DIR / filename)
    else:
        return pages_text, empty_stats

    if not images:
        return pages_text, empty_stats

    stats = {"total": 0, "success": 0, "failed": 0}
    for pos, image_bytes, mime in images[: settings.VISION_MAX_IMAGES]:
        stats["total"] += 1
        try:
            desc = vision_service.describe_image(image_bytes, mime)
        except Exception as e:
            stats["failed"] += 1
            logger.warning("图片描述失败（跳过，不影响入库）：%s", e)
            continue
        if not desc:
            stats["failed"] += 1
            logger.warning("图片描述为空（VL 没看明白这张图），跳过")
            continue
        stats["success"] += 1
        # PDF：拼到对应页；docx：全拼到唯一那页（pages_text[0]）
        target = pos if 0 <= pos < len(pages_text) else 0
        pages_text[target] = f"{pages_text[target]}\n\n[图：{desc}]"

    return pages_text, stats


# 扩展名 → 解析器（Day 20 分发表）。
# 注意：必须用 lambda 包一层，让 parse_pdf/parse_docx/parse_txt 在【调用时】从模块
# 全局查。直接存函数引用 {'.pdf': parse_pdf} 会把原函数对象拷进 dict，
# pytest 的 monkeypatch.setattr(document_service, 'parse_pdf', ...) 就失效了。
_PARSERS = {
    ".pdf": lambda fp: parse_pdf(fp),
    ".docx": lambda fp: parse_docx(fp),
    ".txt": lambda fp: parse_txt(fp),
    ".md": lambda fp: parse_txt(fp),  # MVP 里 md 就是纯文本，直接复用
}


def parse_document(filename: str, file_path: str) -> list[str]:
    """按扩展名分发到具体解析器，返回 list[str]（对 split_text 透明）

    get_extension 已做小写化，所以 .PDF / .Docx 天然命中。未知扩展名抛 ValueError，
    正常流程到不了（上传已 fail-fast），纯属防御兜底——会被 ingest_document_job
    捕获 → 状态变 failed，error_message 就是这句提示。
    """
    ext = get_extension(filename)
    parser = _PARSERS.get(ext)
    if parser is None:
        raise ValueError(
            f"不支持的文件类型：{ext or '（无扩展名）'}，"
            f"仅支持 {' / '.join(sorted(settings.SUPPORTED_EXTENSIONS))}"
        )
    return parser(file_path)


def split_text(pages_text: list[str], chunk_size: int = 500, chunk_overlap: int = 50) -> list[str]:
    """把多页文本切成 chunk

    参数：
        pages_text: 每一页的文本列表，如 ["第一页...", "第二页..."]
        chunk_size: 每个 chunk 最多多少字符
        chunk_overlap: 相邻 chunk 重叠多少字符

    返回：
        chunk 文本列表，如 ["chunk1...", "chunk2...", ...]
    """
    # 1. 先把多页合并成一大段
    full_text = "\n".join(pages_text)

    # 2. 普通正文走 RecursiveCharacterTextSplitter；表格整体保留为独立 chunk，
    #    避免表头/数据行被 500 字切分器拆开。
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
    )
    table_pattern = re.compile(r"\[表格\]\n(.*?)\n\[表格结束\]", re.DOTALL)
    chunks: list[str] = []
    cursor = 0
    for match in table_pattern.finditer(full_text):
        chunks.extend(splitter.split_text(full_text[cursor : match.start()]))
        table = match.group(1).strip()
        if table:
            chunks.append("[表格]\n" + table)
        cursor = match.end()
    chunks.extend(splitter.split_text(full_text[cursor:]))
    return [chunk for chunk in chunks if chunk.strip()]


def split_text_with_pages(
    pages_text: list[str], chunk_size: int = 500, chunk_overlap: int = 50
) -> list[tuple[int, str]]:
    """按页切 chunk，并保留页码，给入库 metadata 用。"""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
    )
    chunks: list[tuple[int, str]] = []
    for page_no, page_text in enumerate(pages_text, start=1):
        if not (page_text or "").strip():
            continue
        table_pattern = re.compile(r"\[表格\]\n(.*?)\n\[表格结束\]", re.DOTALL)
        cursor = 0
        for match in table_pattern.finditer(page_text):
            for chunk in splitter.split_text(page_text[cursor : match.start()]):
                if chunk.strip():
                    chunks.append((page_no, chunk))
            table = match.group(1).strip()
            if table:
                chunks.append((page_no, "[表格]\n" + table))
            cursor = match.end()
        for chunk in splitter.split_text(page_text[cursor:]):
            if chunk.strip():
                chunks.append((page_no, chunk))
    return chunks


def _ingest(db: Session, document_id: int, filename: str, batch_size: int = 32) -> int:
    """把文档变成可检索的向量（解析 → 切分 → 向量化 → 入库）

    这是后台任务真正干的重活。Day 18 改造：
    1. embedding 分批：每批 batch_size 个 chunk 调一次 embed_documents。
       原来把全部 chunk 塞进一个请求，大文件会超 API 的输入上限。
    2. 向量库 + chunk 表的本地写用 _ingest_lock 串行，防并发入库撞车。
    3. 状态流转不归这里管（ingest_document_job 负责），只干"入库"本身。

    返回：切出来的 chunk 数量
    """
    # 1. 按扩展名解析（用绝对路径，不依赖运行时的工作目录）
    abs_path = STORAGE_DIR / filename
    pages_text = parse_document(filename, str(abs_path))

    # Day 21：图片理解增强层——把文档里的图翻译成文字描述拼回文本流。
    # Day 23：失败自动降级（单图跳过）但不再无声——返回 stats，下面打总结日志。
    pages_text, img_stats = enrich_pages_with_images(filename, pages_text)

    # Day 20：空文档（扫描件伪装成 PDF / 空 txt / 空 docx）显式失败，
    # 不静默"processed + 0 chunk"——用户看到 failed + 原因，才知道
    # "传了个读不出字的文件"。（p or "" 防 None：pdfplumber 对扫描页返回 None）
    if not pages_text or all(not (p or "").strip() for p in pages_text):
        raise ValueError("没有可提取的文字，暂不支持 OCR/扫描件，请上传带文本层的文件")

    # 2. 切块（保留页码，财务问答能回到原始页看证据）
    page_chunks = split_text_with_pages(pages_text)

    # 理论兜底（正常会被上面的空检查拦下）：极端情况下 split 出空列表
    if not page_chunks:
        return 0

    # 3. 分批向量化 + 入库
    for i in range(0, len(page_chunks), batch_size):
        batch = page_chunks[i : i + batch_size]
        batch_texts = [chunk for _, chunk in batch]
        ids = [f"doc_{document_id}_chunk_{j}" for j in range(i, i + len(batch))]
        # 给图片和表格 chunk 打类型标记，检索层可以按内容形态过滤/加权。
        metadatas = [
            {
                "source": filename,
                "document_id": document_id,
                "chunk_index": j,
                "page": page_no,
                **(
                    {"type": "image"}
                    if "[图：" in chunk
                    else {"type": "table"} if chunk.startswith("[表格]\n") else {}
                ),
            }
            for (page_no, chunk), j in zip(batch, range(i, i + len(batch)))
        ]

        embeddings = embedding_service.embed_documents(batch_texts)  # 慢的 API 调用，放锁外

        with _ingest_lock:  # 本地快写，串行防撞
            vector_store.add(
                ids=ids,
                embeddings=embeddings,
                documents=batch_texts,
                metadatas=metadatas,
            )
            # 把 chunk 也记到 MySQL（document_chunks 表），便于追溯
            for idx, (cid, (_, chunk)) in enumerate(zip(ids, batch)):
                db.add(
                    DocumentChunk(
                        document_id=document_id,
                        content=chunk,
                        chunk_index=i + idx,
                        vector_id=cid,
                    )
                )
        db.commit()

    # Day 23：丢图可感知——每次入库结束打一行总结，重传文档就能在日志里看到
    # "有几张图没进去、为什么"。之前静默降级，用户根本不知道哪张图丢了。
    logger.info(
        "文档 %s：尝试描述 %d 张图，成功 %d，失败跳过 %d",
        filename, img_stats["total"], img_stats["success"], img_stats["failed"],
    )
    return len(page_chunks)


def ingest_document_job(document_id: int, filename: str) -> None:
    """后台入库任务：自己开会话跑 _ingest，并负责状态流转（Day 18 异步核心）

    为什么不能复用上传请求的 db？
    FastAPI 的 BackgroundTasks 在响应发完之后才跑，而请求的数据库会话（get_db）
    在响应结束时就关了。所以这里必须自己 SessionLocal() 开一个全新会话，
    干完就关（try/finally），绝不泄漏连接。

    status 流转：uploading → processing → processed（成功）
                              ↘ failed + error_message（异常原因存起来，前端可查）
    """
    db = SessionLocal()
    try:
        db.query(Document).filter(Document.id == document_id).update(
            {"status": "processing"}
        )
        db.commit()

        _ingest(db, document_id, filename)

        db.query(Document).filter(Document.id == document_id).update(
            {"status": "processed", "error_message": None}
        )
        db.commit()
    except Exception as e:
        # 任何一步挂了：先回滚失败事务，再标记 failed + 记原因（截断到列宽）
        db.rollback()
        db.query(Document).filter(Document.id == document_id).update(
            {"status": "failed", "error_message": str(e)[:500]}
        )
        db.commit()
    finally:
        db.close()


def delete_document(db: Session, document_id: int) -> Document | None:
    """删除一个文档（Day 14 文档管理）

    顺序很关键：先清"引用"（向量库、chunks 表），再删"主体"（记录、文件）。
    从外向里删，避免删了主体还留着悬空的向量/chunk。

    返回被删的 Document；不存在返回 None。
    """
    doc = db.query(Document).filter(Document.id == document_id).first()
    if doc is None:
        return None

    # 1. 清向量库：按 metadata 里的 document_id 过滤删除
    vector_store.delete_by_document(document_id)

    # 2. 清 chunks 表
    db.query(DocumentChunk).filter(DocumentChunk.document_id == document_id).delete()

    # 3. 删磁盘文件（file_path 是相对 backend 根目录存的，转绝对路径）
    backend_root = Path(__file__).resolve().parent.parent.parent
    abs_path = backend_root / doc.file_path
    if abs_path.exists():
        abs_path.unlink()

    # 4. 删记录
    db.delete(doc)
    db.commit()
    return doc
