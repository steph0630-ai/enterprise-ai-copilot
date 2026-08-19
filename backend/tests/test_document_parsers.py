"""文档解析层测试（Day 20：多格式入库）

两类姿势，原则都是"不碰真实向量库/API"：
1. 真实文件（tmp_path 托管、自动清理）：用具体字节/真实 docx 测具体解析器，
   特别是 TXT 的 GBK 编码坑——这是企业 Windows 文件最典型的雷。
2. monkeypatch 打桩：测 parse_document 分发逻辑。分发测试绝不碰真实文件系统，
   路径随便传，解析器已桩，只要验证"按扩展名找到对的解析器"。
"""

import pytest
from docx import Document as DocxDocument

from app.services import document_service


# ========== 文件名工具函数（路径穿越防线） ==========

def test_sanitize_filename_removes_traversal():
    """正斜杠/反斜杠路径成分都被剥掉，只留文件名（防 ../ 逃出目录）"""
    assert document_service.sanitize_filename("../../../etc/passwd") == "passwd"
    assert document_service.sanitize_filename("..\\..\\report.pdf") == "report.pdf"
    assert document_service.sanitize_filename("C:/x/y/合同.pdf") == "合同.pdf"
    assert document_service.sanitize_filename("") == ""
    assert document_service.sanitize_filename(None) == ""


def test_get_extension_normalizes():
    """扩展名统一小写、带点；无扩展名/空名返回 ''"""
    assert document_service.get_extension("A.PDF") == ".pdf"
    assert document_service.get_extension("x.txt") == ".txt"
    assert document_service.get_extension("report.DocX") == ".docx"
    assert document_service.get_extension("noext") == ""
    assert document_service.get_extension("x.pdf ") == ".pdf"  # 空格被 strip
    assert document_service.get_extension("") == ""


# ========== parse_txt：编码逐级降级 ==========

def test_parse_txt_utf8(tmp_path):
    p = tmp_path / "a.txt"
    p.write_bytes("中文内容 hello".encode("utf-8"))
    assert document_service.parse_txt(str(p)) == ["中文内容 hello"]


def test_parse_txt_gbk_fallback(tmp_path):
    """GBK 编码（Windows 老文件）也能解出中文——Docker 默认 UTF-8 会崩，这里兜住"""
    p = tmp_path / "gbk.txt"
    p.write_bytes("报销流程\n第一步填表".encode("gbk"))
    assert document_service.parse_txt(str(p)) == ["报销流程\n第一步填表"]


def test_parse_txt_utf8_bom(tmp_path):
    """UTF-8 BOM 会被 utf-8-sig 自动剥掉，不残留 \\ufeff"""
    p = tmp_path / "bom.txt"
    p.write_bytes("﻿标题".encode("utf-8"))
    out = document_service.parse_txt(str(p))
    assert out == ["标题"]
    assert "﻿" not in out[0]


def test_parse_txt_empty_file(tmp_path):
    """0 字节空文件 → [""]，交给 _ingest 的空检查判定失败"""
    p = tmp_path / "empty.txt"
    p.write_bytes(b"")
    assert document_service.parse_txt(str(p)) == [""]


def test_parse_md_reads(tmp_path):
    """Markdown 在 MVP 里就是纯文本，直接读"""
    p = tmp_path / "note.md"
    p.write_bytes("# 标题\n\n正文内容".encode("utf-8"))  # bytes 字面量只能含 ASCII，中文要 encode
    assert document_service.parse_txt(str(p)) == ["# 标题\n\n正文内容"]


# ========== parse_docx：段落 + 表格按顺序 ==========

def _make_docx(path, paragraphs=None, table=None):
    """用 python-docx 现场造一份真实 .docx，避免依赖测试夹具文件"""
    d = DocxDocument()
    for para in paragraphs or []:
        d.add_paragraph(para)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, val in enumerate(row):
                t.cell(i, j).text = val
    d.save(path)
    return path


def test_parse_docx_paragraphs(tmp_path):
    p = _make_docx(tmp_path / "para.docx", paragraphs=["第一段", "第二段"])
    out = document_service.parse_docx(str(p))
    assert len(out) == 1  # 整篇作为一"页"
    assert "第一段" in out[0]
    assert "第二段" in out[0]


def test_parse_docx_table(tmp_path):
    p = _make_docx(tmp_path / "table.docx", table=[["部门", "人数"], ["销售一部", "20"]])
    out = document_service.parse_docx(str(p))
    assert "销售一部" in out[0]
    assert " | " in out[0]  # 单元格用竖线分隔


def test_parse_docx_paragraph_table_order(tmp_path):
    """段落-表格-段落的顺序被保留（doc.paragraphs/doc.tables 会丢顺序）

    注意：add_paragraph / add_table 都追加到文档末尾。要造"段落-表格-段落"
    的真实顺序，必须按"段一 → 表格 → 段三"的顺序构造，不能先加完全部段落。
    """
    d = DocxDocument()
    d.add_paragraph("条款一")
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "条款二"
    d.add_paragraph("条款三")
    p = tmp_path / "order.docx"
    d.save(p)

    out = document_service.parse_docx(str(p))[0]
    assert out.index("条款一") < out.index("条款二") < out.index("条款三")


def test_parse_docx_corrupt_raises(tmp_path):
    """假 .docx（实际是文本）→ 中文 ValueError，不是晦涩的 zip 异常"""
    p = tmp_path / "fake.docx"
    p.write_bytes(b"this is not a docx")
    with pytest.raises(ValueError, match="无法解析 Word 文档"):
        document_service.parse_docx(str(p))


# ========== parse_document：按扩展名分发（monkeypatch 桩，不碰真实文件） ==========

def test_parse_document_dispatches_pdf(monkeypatch):
    monkeypatch.setattr(document_service, "parse_pdf", lambda fp: ["pdf"])
    assert document_service.parse_document("a.pdf", "/no/such.pdf") == ["pdf"]


def test_parse_document_dispatches_docx(monkeypatch):
    monkeypatch.setattr(document_service, "parse_docx", lambda fp: ["docx"])
    assert document_service.parse_document("a.docx", "/no/such.docx") == ["docx"]


def test_parse_document_dispatches_txt_md(monkeypatch):
    monkeypatch.setattr(document_service, "parse_txt", lambda fp: ["txt"])
    assert document_service.parse_document("a.txt", "/x") == ["txt"]
    assert document_service.parse_document("a.md", "/x") == ["txt"]


def test_parse_document_uppercase_extension(monkeypatch):
    """A.PDF 大写也命中——get_extension 统一小写了"""
    monkeypatch.setattr(document_service, "parse_pdf", lambda fp: ["pdf"])
    assert document_service.parse_document("A.PDF", "/x") == ["pdf"]


def test_parse_document_unknown_extension_raises(monkeypatch):
    monkeypatch.setattr(document_service, "parse_pdf", lambda fp: ["pdf"])
    with pytest.raises(ValueError, match="不支持的文件类型"):
        document_service.parse_document("a.xlsx", "/x")
    with pytest.raises(ValueError, match="不支持的文件类型"):
        document_service.parse_document("a.png", "/x")


def test_parse_document_no_extension_raises():
    with pytest.raises(ValueError, match="不支持"):
        document_service.parse_document("README", "/x")
